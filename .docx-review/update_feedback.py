from pathlib import Path
from docx import Document
from docx.shared import Pt


source = Path(r"C:\Users\asiro\Downloads\Capstone-Feedback-Verification-G3.docx")
output = Path(r"C:\Users\asiro\Desktop\Capstone\RORMS\Capstone-Feedback-Verification-G3_Updated.docx")
document = Document(source)


def replace_cell_text(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(value)
    run.font.size = Pt(10)


def replace_paragraph_text(paragraph, value: str) -> None:
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(value)
    run.font.size = Pt(10)


project_info = document.tables[0]
replace_cell_text(
    project_info.cell(0, 1),
    "PHINMA University of Pangasinan Room Management System",
)
replace_cell_text(
    project_info.cell(1, 1),
    "A web and mobile-based system for managing academic room assignments and reservations at PHINMA University of Pangasinan. It supports room and building management, prevents schedule conflicts, provides real-time room availability, and generates reports on assigned rooms, schedules, and utilization.",
)

feedback = document.tables[1]
replace_cell_text(
    feedback.cell(1, 1),
    "Change of Title - PHINMA University of Pangasinan Room Management System",
)
replace_cell_text(
    feedback.cell(2, 1),
    "System Scope - Changed from a generic room reservation system to a client-based system for PHINMA University of Pangasinan.",
)
replace_cell_text(
    feedback.cell(3, 1),
    "Room and Building Management - Supports creating buildings with their rooms and floors, as well as adding individual rooms to existing buildings.",
)
replace_cell_text(
    feedback.cell(5, 1),
    "System Functionality - Expanded from a room-reservation-only system into an integrated room-assignment and room-reservation system. The system now assigns rooms to academic schedules and supports reservation requests for other approved campus activities.",
)
replace_cell_text(
    feedback.cell(6, 1),
    "Report Generation - Generates reports for each room showing assigned hours, days, and instructors, and provides each instructor's assigned room and schedule.",
)

for paragraph in document.paragraphs:
    if "Summary of Panelists" in paragraph.text:
        replace_paragraph_text(paragraph, "Summary of Panelists' Feedback and Required Revisions")
    elif "We, the undersigned panelists" in paragraph.text:
        replace_paragraph_text(
            paragraph,
            "We, the undersigned panelists, confirm that the above-listed feedback/revisions were discussed during the capstone project defense and will serve as the basis for the group's revisions for the next semester.",
        )

document.save(output)
print(output)
